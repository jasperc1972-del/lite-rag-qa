"""
test_loaders.py — 讀檔模組與格式註冊表 (F-00b、F-00c、F-01) 的單元測試

測試用的 PDF / Word 檔都在記憶體中即時產生，不需要準備實體檔案。
"""
import io

import pytest

from loaders import (LoaderError, ScannedPdfError, UnsupportedFormatError,
                     docx_loader, get_loader, pdf_loader, supported_extensions,
                     text_loader)


# ---------- 格式註冊表 ----------

def test_R1支援的副檔名():
    assert sorted(supported_extensions()) == [".docx", ".md", ".pdf", ".txt"]


def test_副檔名不分大小寫():
    assert get_loader("規格書.PDF") is pdf_loader.load
    assert get_loader("說明.Docx") is docx_loader.load


@pytest.mark.parametrize("name, release", [("報價.xlsx", "R2"), ("清單.csv", "R2"), ("照片.png", "R3")])
def test_後續版本格式被拒並提示版本(name, release):
    with pytest.raises(UnsupportedFormatError, match=release):
        get_loader(name)


@pytest.mark.parametrize("name", ["程式.exe", "沒有副檔名"])
def test_不在清單內的格式被拒(name):
    with pytest.raises(UnsupportedFormatError):
        get_loader(name)


# ---------- text_loader（.md / .txt）----------

def test_UTF8依空行分段():
    data = "第一段\n第一段第二行\n\n\n第二段\r\n\r\n第三段".encode("utf-8")
    docs = text_loader.load("a.txt", data)
    assert [d.text for d in docs] == ["第一段\n第一段第二行", "第二段", "第三段"]
    assert [d.metadata["paragraph"] for d in docs] == [1, 2, 3]
    assert docs[0].metadata["format"] == "txt"


def test_記事本BOM自動去除():
    docs = text_loader.load("a.md", "﻿標題".encode("utf-8"))
    assert docs[0].text == "標題"


def test_Big5編碼也能讀():
    docs = text_loader.load("舊檔.txt", "繁體中文舊系統".encode("cp950"))
    assert docs[0].text == "繁體中文舊系統"


def test_無法辨識的編碼回報錯誤():
    with pytest.raises(LoaderError, match="編碼"):
        text_loader.load("壞檔.txt", b"\xff\xfe\xfd\xfc")


# ---------- docx_loader（Word）----------

def _make_docx() -> bytes:
    """產生一份含段落、空白段落與表格的 Word 檔"""
    import docx
    word = docx.Document()
    word.add_paragraph("前言")
    word.add_paragraph("")                 # 空白段落應被略過
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "品名", "數量"
    table.cell(1, 0).text, table.cell(1, 1).text = "螺絲 M6", "200"
    word.add_paragraph("結語")
    buf = io.BytesIO()
    word.save(buf)
    return buf.getvalue()


def test_Word依原順序讀出段落與表格():
    docs = docx_loader.load("a.docx", _make_docx())
    assert [d.text for d in docs] == ["前言", "品名｜數量", "螺絲 M6｜200", "結語"]
    assert [d.metadata["paragraph"] for d in docs] == [1, 2, 3, 4]


def test_非Word檔回報錯誤():
    with pytest.raises(LoaderError):
        docx_loader.load("假的.docx", b"not a docx file")


# ---------- pdf_loader ----------

def _make_pdf(pages: list) -> bytes:
    """產生 PDF；pages 中每個元素是一頁的文字，空字串代表空白頁"""
    import pymupdf
    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        if text:
            page.insert_text((72, 72), text)
    return pdf.tobytes()


def test_PDF逐頁讀取並略過空白頁():
    docs = pdf_loader.load("a.pdf", _make_pdf(["Page one", "", "Page three"]))
    assert [d.text for d in docs] == ["Page one", "Page three"]
    assert [d.metadata["page"] for d in docs] == [1, 3]   # 頁碼保留原本位置


def test_沒有文字層判定為掃描檔():
    with pytest.raises(ScannedPdfError, match="R3"):
        pdf_loader.load("掃描.pdf", _make_pdf(["", ""]))


def test_損毀PDF回報錯誤():
    with pytest.raises(LoaderError):
        pdf_loader.load("壞檔.pdf", b"%PDF-broken")
